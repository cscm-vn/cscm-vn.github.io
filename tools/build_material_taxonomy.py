"""Build data/registry/material_taxonomy.json from the team's master-data workbook (Masterise).
Usage: python tools/build_material_taxonomy.py path/to/master.xlsx
The workbook itself is NOT committed (contains BOQ values / internal scoring). Only names, levels,
codes, units and sample work-detail wording are exported. Codes are slugs of names → stable across reruns."""
import sys, json, re, unicodedata as U, collections as C, datetime, os
import openpyxl
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
def nfc(s): return U.normalize("NFC", str(s).strip()) if s is not None else ""
def slug(s):
    s = "".join(c for c in U.normalize("NFD", s) if U.category(c) != "Mn").replace("đ", "d").replace("Đ", "D")
    s = s.replace("≥", "ge").replace("<", "lt").replace("&", "-")
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")[:40]
wb = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
nodes = {}
def add(code, parent, level, name, src, **kw):
    if code not in nodes:
        nodes[code] = dict(code=code, parent=parent, level=level, name_vi=nfc(name), description="", source_sheet=src, **kw)
    return code
L1 = {"KC": ("Kết cấu & phần thô", "2.1 Ma_tran_QD_Ket_cau", "d_masterdata_KC"),
      "MEPF": ("Cơ điện (MEPF)", "2.2 Ma_tran_QD_MEPF", "d_masterdata_MEPF"),
      "HT": ("Hoàn thiện", "2.3 Ma_tran_QD_Hoàn thiện", "d_masterdata_Hoàn thiện")}
HE = {"Điện": "Điện", "HVAC": "Điều hòa & thông gió (HVAC)", "Nước": "Cấp thoát nước", "PCCC": "Phòng cháy chữa cháy (PCCC)", "ELV": "Điện nhẹ (ELV)", "Khác": "Khác"}
samples, units, supply, rows_n = C.defaultdict(list), C.defaultdict(C.Counter), {}, C.Counter()
for k, (name, msheet, dsheet) in L1.items():
    add(k, None, 1, name, msheet)
    # matrix: price families + hệ + current supplier status
    fam_he = {}
    for r in list(wb[msheet].iter_rows(values_only=True))[2:]:
        if k == "HT":
            fam, he = r[1], None   # HT matrix col B = price family (XLOOKUP), col C = level-2 family
        else:
            fam, he = r[1], r[2]
        if fam and r[0]:
            fam_he[nfc(fam)] = nfc(he) if he else None
            supply[(k, nfc(fam) if k != "HT" else nfc(r[2]))] = nfc(r[0])
    for r in list(wb[dsheet].iter_rows(values_only=True))[1:]:
        if not r[0]: continue
        fam, fam2, sysl1 = nfc(r[9]), nfc(r[10]), nfc(r[2])
        if not fam or fam.startswith("="): fam = fam2
        if k == "KC":
            p = add(f"KC.{slug(fam)}", "KC", 2, fam, dsheet)
        elif k == "MEPF":
            he = fam_he.get(fam) or "Khác"
            h = add(f"MEPF.{slug(he)}", "MEPF", 2, HE.get(he, he), msheet, short=he)
            p = add(f"{h}.{slug(fam)}", h, 3, fam, dsheet)
            if fam2 and fam2 != fam:
                p = add(f"{p}.{slug(fam2)}", p, 4, fam2, dsheet)
        else:
            s = add(f"HT.{slug(sysl1)}", "HT", 2, sysl1.capitalize() if False else sysl1, dsheet)
            disp = re.sub(r"\s*-\s*(CC&LĐ|Cung cấp|D&B|Thi công)$", "", fam)
            f = add(f"{s}.{slug(disp)}", s, 3, disp, dsheet, package=fam)
            p = add(f"{f}.{slug(fam2)}", f, 4, fam2, dsheet)
        wd = nfc(r[5]) or nfc(r[4])
        if wd and wd not in samples[p] and len(samples[p]) < 4: samples[p].append(wd)
        if r[6]: units[p][nfc(r[6])] += 1
        x = p
        while x: rows_n[x] += 1; x = nodes[x]["parent"]
for c, n in nodes.items():
    n["master_rows"] = rows_n[c]
    if samples[c]: n["description"] = "VD trong master data: " + "; ".join(samples[c])
    if units[c]: n["units"] = [u for u, _ in units[c].most_common(3)]
    top = c.split(".")[0]
    st = supply.get((top, n["name_vi"]))
    if st: n["supply_now"] = st
# KC families mentioned as moved to HT ('→HT') keep their wording; flag them
for n in nodes.values():
    if "(→HT)" in n["name_vi"]: n["note"] = "Master data ghi '→HT': đánh giá ở ma trận Hoàn thiện"
out = dict(schema_version=1, name="Cây nhóm vật tư theo master data (ngôn ngữ nhóm Mua sắm)",
           source=dict(file="Master data ma trận quyết định (xlsx, nội bộ – không commit)", sheets=[v[1] for v in L1.values()] + [v[2] for v in L1.values()],
                       built=datetime.date.today().isoformat(), excluded="Điểm số, BOQ/giá trị, remark chiến lược không xuất ra web công khai"),
           levels={"1": "Khối (sheet ma trận)", "2": "Hệ / System (Level 1)", "3": "Nhóm giá – Price family", "4": "Nhóm giá Level 2"},
           note="KC: Level 2 = Nhóm giá (chỉ có 1 hệ 'Kết cấu'). MEPF: L2 = Hệ (cột 'Hệ' ma trận 2.2), L3 = Nhóm giá, L4 = Nhóm giá Level 2 nếu khác. HT: L2 = System (Level 1), L3 = gói (Price family, bỏ hậu tố CC&LĐ/Cung cấp/D&B), L4 = Price family Level 2.",
           nodes=sorted(nodes.values(), key=lambda n: n["code"]))
json.dump(out, open(os.path.join(ROOT, "data/registry/material_taxonomy.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
print(len(nodes), C.Counter(n["level"] for n in nodes.values()))
