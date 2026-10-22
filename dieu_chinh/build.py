"""Điều chỉnh giá HĐ tổng thầu: so giá thị trường (data/history/prices) với đơn giá nguyên tắc CTD/NTC (data/registry/contract_baselines.json).
Chạy SAU export_static.py và forecast/build.py, TRƯỚC trends/build.py:  python dieu_chinh/build.py --site site
→ site/data/dieu-chinh.json, site/dieu-chinh.csv, site/dieu-chinh.xlsx. Không bịa giá: thiếu nguồn → 'missing'."""
import argparse, csv, json, math, os, statistics, unicodedata, datetime as dt
from zoneinfo import ZoneInfo
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TZ = ZoneInfo("Asia/Ho_Chi_Minh")
J = lambda p, d=None: json.load(open(p, encoding="utf-8")) if os.path.exists(p) else d
NFC = lambda s: unicodedata.normalize("NFC", s) if isinstance(s, str) else s
RNAME = {"bac": "Miền Bắc", "nam": "Miền Nam", "trung": "Miền Trung"}; FCREG = {"bac": "HN", "nam": "HCM"}
STATUS_VI = {"down": "Vượt ngưỡng GIẢM", "up": "Vượt ngưỡng TĂNG", "within": "Trong ngưỡng", "missing": "Thiếu dữ liệu", "na": "Không áp dụng"}
def Phi(x): return 0.5 * (1 + math.erf(x / math.sqrt(2)))

def load_inputs(site):
    prices = {x["id"]: x for x in J(os.path.join(site, "prices.json"), {"items": []})["items"]}
    hist = {}
    def H(i):
        if i not in hist:
            h = J(os.path.join(site, "history", i + ".json"), {"points": []})
            hist[i] = sorted([p for p in h.get("points", []) if p.get("price") is not None], key=lambda p: p["date"])
        return hist[i]
    fc = {x["id"]: x for x in (J(os.path.join(site, "forecast.json"), {}) or {}).get("items", []) if x.get("fc")}
    return prices, H, fc

def sval(s, p, vat):
    if p is None: return None
    v = p.get("price")
    if v is None: return None
    if s.get("use") == "mid": v = (v + (p.get("price_high") or v)) / 2
    elif s.get("use") == "high": v = p.get("price_high") or v
    v *= s.get("factor", 1.0)
    if s.get("vat") == "incl": v /= (1 + vat)
    return v + s.get("adj", 0)

def series_of(s, H, vat):
    return [(p["date"], sval(s, p, vat)) for p in H(s["id"]) if sval(s, p, vat) is not None]

def combine(list_of_series, how=statistics.median, start=None):
    """carry-forward align on union of dates; value = how(available)."""
    dates = sorted({d for S in list_of_series for d, _ in S})
    if start: dates = [d for d in dates if d >= start]
    out, idx, last = [], [0] * len(list_of_series), [None] * len(list_of_series)
    for d in dates:
        for k, S in enumerate(list_of_series):
            while idx[k] < len(S) and S[idx[k]][0] <= d: last[k] = S[idx[k]][1]; idx[k] += 1
        vals = [v for v in last if v is not None]
        if vals and len(vals) == len(list_of_series) or (vals and how is statistics.median): out.append((d, how(vals)))
    return out

VATL = {"excl": "chưa VAT", "incl": "đã gồm VAT (quy về chưa VAT)", "unknown": "nguồn không nêu VAT (coi là chưa VAT)", "mixed": "nhiều nguồn – đã quy về chưa VAT"}
def basis_text(mk, r):
    b = mk.get("basis")
    if mk.get("value") is None or not b: return None
    src_vat = {"excl": "nguồn chưa VAT", "incl": "nguồn đã gồm VAT → ÷1,08", "unknown": "nguồn không nêu VAT → coi là chưa VAT", "mixed": "nhiều nguồn, mỗi nguồn quy về chưa VAT"}[b.get("vat", "unknown")]
    parts = [f"Giá thị trường quy về CHƯA VAT ({src_vat})"]
    if b.get("delivery"): parts.append("Điều kiện giao: " + "; ".join(b["delivery"]))
    if b.get("adj"): parts.append("Điều chỉnh quy cách: " + "; ".join(b["adj"]))
    if b.get("conv"): parts.append(f"Quy đổi về đ/{r['unit']}: " + "; ".join(b["conv"]))
    return " · ".join(parts)

def build_contract(c, MAP, prices, H, FC, today):
    vat, near_pp, infl = MAP["vat_rate"], MAP["near_pp"], MAP["infl"]
    items = {(r["region"], r["code"]): r for r in c["items"]}
    res = {}  # (region, code) -> market dict + series
    start = (today - dt.timedelta(days=400)).isoformat()
    def mk_median(key, m):
        srcs, cur, main = [], [], []
        for s in m["series"]:
            it = prices.get(s["id"]); S = series_of(s, H, vat)
            if not it or it.get("price") is None: srcs.append(dict(id=s["id"], label=s.get("label"), missing=True)); continue
            v = sval(s, it, vat); cur.append(v)
            if s.get("role") != "check": main.append(v)
            srcs.append(dict(id=s["id"], label=s.get("label"), value=v, role=s.get("role"), raw=it.get("price"), raw_high=it.get("price_high"), unit=it.get("unit"),
                             as_of=it.get("as_of"), src=it.get("src"), url=it.get("url"), stale=bool(it.get("stale")), note=it.get("note"),
                             vat=s.get("vat", "unknown"), vat_ev=s.get("vat_ev"), delivery=s.get("delivery"), adj_note=s.get("adj_note"),
                             conv=" ".join(x for x in [("điểm giữa khoảng" if s.get("use") == "mid" else ""), (f"× {s['factor']:g}".replace(".", ",") if s.get("factor", 1) != 1 else ""),
                                                        ("÷ 1,08 (bỏ VAT)" if s.get("vat") == "incl" else ""),
                                                        (f"+ {s['adj']:,.0f}".replace(",", ".") + " (độ sụt 12±2)" if s.get("adj") else "")] if x) or None))
        if not main: return None
        ser = combine([series_of(s, H, vat) for s in m["series"] if prices.get(s["id"]) and s.get("role") != "check"], start=start)
        ok = [x for x in srcs if not x.get("missing") and x.get("role") != "check"]
        vs = sorted({x["vat"] for x in ok})
        basis = dict(vat=vs[0] if len(vs) == 1 else "mixed", vat_src=[x["vat"] for x in ok], delivery=sorted({x["delivery"] for x in ok if x.get("delivery")}),
                     conv=sorted({x["conv"] for x in ok if x.get("conv")}), adj=sorted({x["adj_note"] for x in ok if x.get("adj_note")}), ev=sorted({x["vat_ev"] for x in ok if x.get("vat_ev")}))
        return dict(value=statistics.median(main), lo=min(cur), hi=max(cur), n=len(cur), n_main=len(main), basis=basis, as_of=min(x["as_of"] for x in ok if x.get("as_of")),
                    newest=max(x["as_of"] for x in ok if x.get("as_of")), stale=any(x["stale"] for x in ok), sources=srcs, series=ser)
    order = sorted(MAP["items"].items(), key=lambda kv: kv[1]["method"] != "median")
    got = lambda k: (lambda x: x if x and x.get("value") is not None else None)(res.get(k))
    for key, m in order:
        reg, code = key.split("|", 1)
        if (reg, code) not in items: continue
        out = None
        if m["method"] == "median": out = mk_median(key, m)
        elif m["method"] == "ratio":
            rr, rc = m["ref"].split("|"); ref = got((rr, rc)); a, b = items[(reg, code)], items[(rr, rc)]
            k = (a["ctd"] / b["ctd"]) if a.get("ctd") and b.get("ctd") else ((a["ntc"] / b["ntc"]) if a.get("ntc") and b.get("ntc") else None)
            if ref and k:
                out = dict(value=ref["value"] * k, lo=ref["lo"] * k, hi=ref["hi"] * k, n=ref["n"], as_of=ref["as_of"], newest=ref["newest"], stale=ref["stale"], basis=dict(ref.get("basis") or {}, conv=[f"giá {rc} × {k:.4f}".replace(".", ",")]),
                           sources=[dict(derived=True, label=f"Giá thị trường {rc} × {k:.4f}".replace(".", ","), value=ref["value"] * k)], series=[(d, v * k) for d, v in ref["series"]])
        elif m["method"] == "interp":
            refs = [got(tuple(x.split("|"))) for x in m["refs"]]
            if all(refs):
                ser = combine([r["series"] for r in refs], how=lambda v: sum(v) / len(v), start=start)
                out = dict(value=sum(r["value"] for r in refs) / len(refs), lo=sum(r["lo"] for r in refs) / len(refs), hi=sum(r["hi"] for r in refs) / len(refs), basis=dict(refs[0].get("basis") or {}, conv=["TB(" + ", ".join(x.split("|")[1] for x in m["refs"]) + ")"]),
                           n=min(r["n"] for r in refs), as_of=min(r["as_of"] for r in refs), newest=max(r["newest"] for r in refs), stale=any(r["stale"] for r in refs),
                           sources=[dict(derived=True, label=f"TB({', '.join(x.split('|')[1] for x in m['refs'])})", value=sum(r["value"] for r in refs) / len(refs))], series=ser)
        elif m["method"] == "norm":
            xm, ca = got((reg, "XM")), got((reg, "CATTO")); bx, bc = items.get((reg, "XM"), {}).get("ntc"), items.get((reg, "CATTO"), {}).get("ntc")
            kg, m3 = m["inputs"]["XM"], m["inputs"]["CATTO"]; base = items[(reg, code)].get("ntc") or items[(reg, code)].get("ctd")
            if xm and ca and bx and bc and base:
                c0 = kg * bx + m3 * bc
                f = lambda x, y: base * (kg * x + m3 * y) / c0
                ser = [(d, base * v / c0) for d, v in combine([[(d, kg * v) for d, v in xm["series"]], [(d, m3 * v) for d, v in ca["series"]]], how=sum, start=start)]
                bx_, bc_ = xm.get("basis") or {}, ca.get("basis") or {}
                out = dict(value=f(xm["value"], ca["value"]), lo=f(xm["lo"], ca["lo"]), hi=f(xm["hi"], ca["hi"]), n=min(xm["n"], ca["n"]),
                           basis=dict(vat=bx_.get("vat") if bx_.get("vat") == bc_.get("vat") else "mixed", delivery=bx_.get("delivery", []) + bc_.get("delivery", []),
                                      conv=[f"đơn giá NTC × chỉ số chi phí ({kg} kg XM + {m3:g} m³ cát)".replace(".", ",")], adj=[], ev=bx_.get("ev", []) + bc_.get("ev", [])),
                           as_of=min(xm["as_of"], ca["as_of"]), newest=max(xm["newest"], ca["newest"]), stale=xm["stale"] or ca["stale"],
                           sources=[dict(derived=True, label=f"{kg} kg XM × {xm['value']:,.0f} + {m3:g} m³ cát × {ca['value']:,.0f}".replace(",", ".") + f" = {kg * xm['value'] + m3 * ca['value']:,.0f}".replace(",", ".") +
                                         f" so với {c0:,.0f}".replace(",", ".") + " theo đơn giá NTC", value=f(xm["value"], ca["value"]))], series=ser,
                           shares={"cement": kg * xm["value"] / (kg * xm["value"] + m3 * ca["value"])})
        if out:
            out.update(method=m["method"], quality=m.get("quality"), note=m.get("note"))
            res[(reg, code)] = out
        else:
            refs = []
            for s in m.get("refs_info", []) + (m.get("series", []) if m["method"] == "median" else []):
                it = prices.get(s["id"])
                if it: refs.append(dict(id=s["id"], label=s.get("label"), value=it.get("price"), unit=it.get("unit"), as_of=it.get("as_of"), src=it.get("src"), url=it.get("url")))
            res[(reg, code)] = dict(value=None, method=m["method"], note=m.get("note") or "Chưa có nguồn", refs=refs, why=("Nguồn lỗi/không có giá" if m["method"] != "none" else None))
    # forecasts
    def fcp(fid, h):
        x = FC.get(fid)
        if not x: return None
        f = x["fc"][str(h)]; return f["pct"] / 100, math.log(f["hi"] / f["lo"]) / (2 * 1.2816)
    prof, byc, byg = MAP["forecast_profiles"], MAP["forecast_by_code"], MAP["forecast_by_group"]
    def cement_step(reg, h, now_per_kg, as_of):
        p = prof["cement"]; step = p["step_vnd_per_t"] / ((1 + vat) if p["step_vat_incl"] else 1) / 1000
        return step / now_per_kg if as_of and as_of < p["step_date"] else 0.0
    def forecast(reg, code, grp, mk):
        if not mk or mk.get("value") is None: return None
        R = FCREG.get(reg); kind = byc.get(code) or byg.get(grp)
        if not kind: return None
        out = {}
        for h in (3, 6, 12):
            if kind == "concrete":
                b = fcp(f"rmc400_{R}", h)
                if not b: return None
                g = code[1:]; kg = prof["concrete"]["cement_kg_per_m3"].get(g, 450)
                step = kg / 1000 * 100000 / (1 + vat) / mk["value"]
                pct, sd = b[0] + step, b[1]; note = prof["concrete"]["note"]
            elif kind == "cement":
                b = fcp(f"cement_{R}", h)
                if not b: return None
                pct, sd = b[0] + cement_step(reg, h, mk["value"], mk.get("as_of")), b[1]; note = prof["cement"]["note"]
            elif kind == "sand":
                b = fcp(f"sand_{R}", h)
                if not b: return None
                pct, sd = b; note = prof["sand"]["note"]
            elif kind == "norm":
                xm = res.get((reg, "XM")); bc, bs = fcp(f"cement_{R}", h), fcp(f"sand_{R}", h)
                if not (bc and bs and xm): return None
                w = mk.get("shares", {}).get("cement", 0.4); cst = bc[0] + cement_step(reg, h, xm["value"], xm.get("as_of"))
                pct, sd = w * cst + (1 - w) * bs[0], math.hypot(w * bc[1], (1 - w) * bs[1])
                note = f"Theo cấu thành định mức: {w * 100:.0f}% xi măng (mô hình xi măng + bước tăng 10/2026) + {(1 - w) * 100:.0f}% cát (mô hình cát)."
            else:
                P = prof[kind]; pct, var, sw = 0.0, 0.0, 0.0; missing = []
                for k, wt in P["proxy"].items():
                    fid = {"cement": f"cement_{R}", "sand": f"sand_{R}"}.get(k, k); b = fcp(fid, h)
                    if not b: missing.append(k); continue
                    extra = cement_step(reg, h, (res.get((reg, "XM")) or {}).get("value") or 1500, (res.get((reg, "XM")) or {}).get("as_of")) if k == "cement" else 0
                    pct += wt * (b[0] + extra); var += (wt * b[1]) ** 2; sw += wt
                pct += (1 - sw) * infl * h / 12; sd = math.sqrt(var + (0.06 * math.sqrt(h / 12)) ** 2)
                note = P["note"] + (f" Thiếu dự báo {', '.join(missing)} → phần này tính theo lạm phát." if missing else "")
            now = mk["value"]; mu = math.log(now * (1 + pct)); e = dict(pct=round(pct * 100, 2), base=now * (1 + pct), lo=now * (1 + pct) * math.exp(-1.2816 * sd), hi=now * (1 + pct) * math.exp(1.2816 * sd))
            it = items[(reg, code)]
            for b in ("ctd", "ntc"):
                thr = it.get(f"thr_{b}_pct")
                for suf, kb in (("", 1.0), ("_vi", 1 / (1 + vat))):
                    bv = it.get(b) and it[b] * kb; mu_ = mu + (math.log(kb) if kind == "norm" else 0)
                    if bv and thr is not None and sd > 0:
                        e[f"p_up_{b}{suf}"] = round(100 * (1 - Phi((math.log(bv * (1 + thr / 100)) - mu_) / sd)), 1)
                        e[f"p_down_{b}{suf}"] = round(100 * Phi((math.log(bv * (1 - thr / 100)) - mu_) / sd), 1)
            out[str(h)] = e
        return dict(h=out, note=note, kind=kind)
    rows = []
    for r in c["items"]:
        mk = res.get((r["region"], r["code"])) or dict(value=None, method="none", note="Chưa khai báo ánh xạ giá thị trường")
        evals = {}
        for scen, kb in (("excl", 1.0), ("incl", 1 / (1 + vat))):
          ev = evals[scen] = {}
          for b in ("ctd", "ntc"):
            bv, thr = r.get(b) and r[b] * kb, r.get(f"thr_{b}_pct")
            if not bv: ev[b] = dict(status="na"); continue
            if mk.get("value") is None: ev[b] = dict(status="missing"); continue
            mv = mk["value"] * (kb if mk.get("method") == "norm" else 1)  # 'norm' suy từ đơn giá NTC → cùng cơ sở VAT với HĐ
            d = (mv / bv - 1) * 100
            st = "down" if d < -thr else "up" if d > thr else "within"
            e = dict(dev=round(d, 2), status=st, near=bool(st == "within" and thr - abs(d) <= near_pp), beyond=round((abs(d) - thr) * (1 if d > 0 else -1), 2) if st != "within" else 0.0,
                     gap_pp=round(thr - abs(d), 2), dev_lo=round((mk["lo"] * mv / mk["value"] / bv - 1) * 100, 2), dev_hi=round((mk["hi"] * mv / mk["value"] / bv - 1) * 100, 2))
            if r.get("weight_pct"):
                e["impact"] = round(r["weight_pct"] * d / 100, 3); e["impact_beyond"] = round(r["weight_pct"] * e["beyond"] / 100, 3)
            if (mk.get("basis") or {}).get("vat") == "unknown":  # nguồn không nêu VAT: nếu thực ra đã gồm VAT
                e["dev_if_src_incl"] = round((mv / (1 + vat) / bv - 1) * 100, 2)
            ev[b] = e
        ev = evals["excl"]
        ser = mk.pop("series", None) or []
        monthly = {}
        for d, v in ser: monthly[d[:7]] = v
        row = dict(r, region_name=RNAME.get(r["region"], r["region"]), market={k: v for k, v in mk.items() if k != "shares"}, eval=ev, evals=evals, basis_text=basis_text(mk, r),
                   series=dict(d=[[d, round(v, 2)] for d, v in ser], m=[[k, round(v, 2)] for k, v in sorted(monthly.items())]),
                   fc=forecast(r["region"], r["code"], r["group"], dict(mk, series=None)))
        rows.append(row)
    summs = {}
    for scen in ("excl", "incl"):
      summ = summs[scen] = {}
      for reg in sorted({r["region"] for r in rows}) + ["all"]:
        summ[reg] = {}
        for b in ("ctd", "ntc"):
            E_ = lambda x: x["evals"][scen][b]
            sel = [x for x in rows if reg in ("all", x["region"]) and E_(x)["status"] != "na"]
            S = lambda st: [x for x in sel if E_(x)["status"] == st]
            summ[reg][b] = dict(down=len(S("down")), up=len(S("up")), within=len(S("within")), missing=len(S("missing")), near=sum(1 for x in sel if E_(x).get("near")), n=len(sel),
                                down_w=round(sum(x.get("weight_pct") or 0 for x in S("down")), 2), up_w=round(sum(x.get("weight_pct") or 0 for x in S("up")), 2),
                                impact_down=round(sum(E_(x).get("impact") or 0 for x in S("down")), 3), impact_up=round(sum(E_(x).get("impact") or 0 for x in S("up")), 3),
                                impact_down_beyond=round(sum(E_(x).get("impact_beyond") or 0 for x in S("down")), 3))
    meta = {k: v for k, v in c.items() if k != "items"}
    return dict(meta, rows=rows, summary=summs["excl"], summaries=summs, contract_vat_default=c.get("vat_basis") or MAP.get("contract_vat_default", "excl"), vat_rate=vat,
                vat_note=MAP.get("basis_vat_rate_note"))

def export_tables(data, site):
    hdr = ["Hợp đồng", "Miền", "Nhóm", "Mã", "Mô tả", "ĐVT", "Tỷ trọng (%)", "Đơn giá CTD", "Đơn giá NTC", "Giá thị trường", "Khoảng giá nguồn", "Ngày giá", "Nguồn",
           "Cách xác định", "Độ tin cậy", "Cơ sở so sánh", "% lệch vs CTD", "Ngưỡng CTD (±%)", "Trạng thái CTD", "% lệch vs NTC", "Ngưỡng NTC (±%)", "Trạng thái NTC",
           "Tác động vs CTD (điểm % giá trị HĐ)", "Tác động vs NTC (điểm % giá trị HĐ)", "% lệch vs CTD nếu đơn giá HĐ đã gồm VAT", "Trạng thái CTD nếu HĐ gồm VAT", "% lệch vs NTC nếu đơn giá HĐ đã gồm VAT", "Trạng thái NTC nếu HĐ gồm VAT", "Dự báo 3T (%)", "Dự báo 6T (%)", "Dự báo 12T (%)", "Ghi chú"]
    lines = []
    for c in data["contracts"]:
        for r in c["rows"]:
            m, e = r["market"], r["eval"]
            st = lambda b: STATUS_VI[e[b]["status"]] + (" · sắp chạm ngưỡng" if e[b].get("near") else "")
            src = "; ".join(f"{s.get('label') or s.get('src')}" for s in m.get("sources", []) if not s.get("missing")) if m.get("value") is not None else ""
            fch = (r.get("fc") or {}).get("h", {})
            lines.append([c["name"] + " v" + c["version"], r["region_name"], r["group"], r["code"], r["desc"], r["unit"], r.get("weight_pct"), r.get("ctd"), r.get("ntc"),
                          round(m["value"]) if m.get("value") is not None and m["value"] >= 100 else (round(m["value"], 1) if m.get("value") is not None else None),
                          (f"{m['lo']:,.0f} – {m['hi']:,.0f}".replace(",", ".") if m.get("value") is not None and m.get("n", 1) > 1 else ""), m.get("as_of") or "", src,
                          {"median": "Trung vị nguồn", "ratio": "Suy diễn (tỷ lệ mác)", "interp": "Suy diễn (nội suy)", "norm": "Suy diễn (định mức)", "none": "Thiếu dữ liệu"}.get(m.get("method"), ""),
                          m.get("quality") or "", r.get("basis_text") or "", e["ctd"].get("dev"), r.get("thr_ctd_pct"), st("ctd"), e["ntc"].get("dev"), r.get("thr_ntc_pct"), st("ntc"),
                          e["ctd"].get("impact"), e["ntc"].get("impact"), r["evals"]["incl"]["ctd"].get("dev"), STATUS_VI[r["evals"]["incl"]["ctd"]["status"]],
                          r["evals"]["incl"]["ntc"].get("dev"), STATUS_VI[r["evals"]["incl"]["ntc"]["status"]], (fch.get("3") or {}).get("pct"), (fch.get("6") or {}).get("pct"), (fch.get("12") or {}).get("pct"), m.get("note") or ""])
    with open(os.path.join(site, "dieu-chinh.csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f); w.writerow(hdr); w.writerows(lines)
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook(); ws = wb.active; ws.title = "Theo dõi điều chỉnh giá"
        ws.append([f"Theo dõi điều chỉnh giá HĐ – cập nhật {data['generated']} (giờ VN). Giá thị trường tham khảo nguồn công khai; xem cột Cách xác định / Ghi chú."]); ws["A1"].font = Font(bold=True, size=12, color="14325A")
        ws.append(hdr)
        for cell in ws[2]: cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="14325A"); cell.alignment = Alignment(wrap_text=True, vertical="top")
        fills = {"Vượt ngưỡng GIẢM": "D1FADF", "Vượt ngưỡng TĂNG": "FEE4E2", "Thiếu dữ liệu": "F2F4F7"}
        for ln in lines:
            ws.append(ln)
            for col in (19, 22, 26, 28):
                v = ws.cell(ws.max_row, col).value or ""
                for k, fc in fills.items():
                    if str(v).startswith(k): ws.cell(ws.max_row, col).fill = PatternFill("solid", fgColor=fc)
        from openpyxl.utils import get_column_letter as GL
        for i, wd in enumerate([22, 10, 16, 13, 40, 7, 9, 12, 12, 13, 22, 11, 40, 18, 8, 50, 10, 9, 22, 10, 9, 22, 12, 12, 12, 22, 12, 22, 9, 9, 9, 60], 1):
            ws.column_dimensions[GL(i)].width = wd
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                if isinstance(cell.value, float) and cell.column in (8, 9, 10): cell.number_format = "#,##0"
        ws.freeze_panes = "F3"; ws.auto_filter.ref = f"A2:{GL(len(hdr))}{ws.max_row}"
        wb.save(os.path.join(site, "dieu-chinh.xlsx"))
    except ImportError: pass

def main():
    ap = argparse.ArgumentParser(); ap.add_argument("--site", default=os.path.join(ROOT, "site")); a = ap.parse_args()
    today = dt.datetime.now(TZ).date()
    REGI = J(os.path.join(ROOT, "data", "registry", "contract_baselines.json")); MAP = J(os.path.join(ROOT, "data", "registry", "contract_market_map.json"))
    prices, H, FC = load_inputs(a.site)
    latest = {}
    for c in REGI["contracts"]:
        if c.get("effective_date", "") <= today.isoformat() and (c["id"] not in latest or c["version"] > latest[c["id"]]["version"]): latest[c["id"]] = c
    data = dict(generated=dt.datetime.now(TZ).isoformat(timespec="minutes"), schema_version=1, status_vi=STATUS_VI, near_pp=MAP["near_pp"],
                map_desc=MAP["desc"], forecast_generated=(J(os.path.join(a.site, "forecast.json"), {}) or {}).get("generated"),
                contracts=[build_contract(c, MAP, prices, H, FC, today) for c in latest.values()])
    os.makedirs(os.path.join(a.site, "data"), exist_ok=True)
    s = json.dumps(data, ensure_ascii=False, separators=(",", ":"), default=lambda o: None)
    open(os.path.join(a.site, "data", "dieu-chinh.json"), "w", encoding="utf-8").write(NFC(s))
    export_tables(data, a.site)
    summ = []
    for c in data["contracts"]:
        top = []
        for r in c["rows"]:
            hits = [dict(base=b.upper(), dev=r["eval"][b]["dev"], status=r["eval"][b]["status"]) for b in ("ctd", "ntc") if r["eval"][b]["status"] in ("down", "up")]
            if hits:
                top.append(dict(code=r["code"], region=r["region"], region_name=r["region_name"], hits=hits, status=hits[0]["status"],
                                weight=r.get("weight_pct"), quality=r["market"].get("quality")))
        top.sort(key=lambda x: (x["status"] != "down", -(x["weight"] or 0), min(h["dev"] for h in x["hits"])))
        summ.append(dict(id=c["id"], name=c["name"], version=c["version"], summary=c["summary"], summary_if_contract_incl_vat=c["summaries"]["incl"], vat_basis=c.get("vat_basis"), vat_basis_note=c.get("vat_basis_note"), top=top[:8],
                         basis="Giá thị trường quy về chưa VAT; đơn giá HĐ chưa gồm VAT" + (" (đã xác nhận)" if c.get("vat_basis") == "excl" else " (giả định, chưa xác nhận)")))
    open(os.path.join(a.site, "data", "dieu-chinh-summary.json"), "w", encoding="utf-8").write(NFC(json.dumps(dict(generated=data["generated"], contracts=summ), ensure_ascii=False, separators=(",", ":"))))
    for c in data["contracts"]:
        S = c["summary"]["all"]; print(f"{c['id']} v{c['version']}: {len(c['rows'])} dòng | CTD ↓{S['ctd']['down']} ↑{S['ctd']['up']} ≈{S['ctd']['within']} ?{S['ctd']['missing']} | NTC ↓{S['ntc']['down']} ↑{S['ntc']['up']} ≈{S['ntc']['within']} ?{S['ntc']['missing']}")
if __name__ == "__main__": main()
